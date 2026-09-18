import csv

import openpyxl
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from app.main import app

client = TestClient(app)


def _new_project() -> int:
    resp = client.post("/api/projects", json={"name": "Parsing Test Project"})
    assert resp.status_code == 200
    return resp.json()["id"]


def _upload(project_id: int, path, filename: str, content_type: str, category="spec"):
    with open(path, "rb") as f:
        resp = client.post(
            f"/api/projects/{project_id}/documents",
            files={"file": (filename, f, content_type)},
            data={"category": category},
        )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_pdf_parsing(tmp_path):
    pdf_path = tmp_path / "Motor_Spec.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=(300, 300))
    c.drawString(20, 200, "TAG: M-201")
    c.drawString(20, 180, "Power: 30 kW")
    c.drawString(20, 160, "Voltage: 400 V")
    c.showPage()
    c.save()

    project_id = _new_project()
    doc = _upload(project_id, pdf_path, "Motor_Spec.pdf", "application/pdf")
    assert doc["status"] in {"parsed", "analyzed"}  # "analyzed" if GEMINI_API_KEY happens to be configured

    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert len(chunks) == 1
    assert chunks[0]["location"] == "p.1"
    assert "M-201" in chunks[0]["text"]
    assert "30 kW" in chunks[0]["text"]


def test_xlsx_parsing(tmp_path):
    xlsx_path = tmp_path / "Equipment_List.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["TAG", "Manufacturer", "Power"])
    ws.append(["M-201", "Siemens", 30])
    wb.save(xlsx_path)

    project_id = _new_project()
    doc = _upload(
        project_id, xlsx_path, "Equipment_List.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        category="equipment_list",
    )
    assert doc["status"] in {"parsed", "analyzed"}  # "analyzed" if GEMINI_API_KEY happens to be configured

    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert len(chunks) == 1  # one data row (the header row becomes column context, not its own chunk)
    assert chunks[0]["location"] == "Row2"  # exact physical row — precise enough to trace/highlight
    assert "Siemens" in chunks[0]["text"]
    assert "Row2" in chunks[0]["text"]


def test_csv_parsing(tmp_path):
    csv_path = tmp_path / "Load_Calculation.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["TAG", "Power"])
        writer.writerow(["M-201", "33"])

    project_id = _new_project()
    doc = _upload(project_id, csv_path, "Load_Calculation.csv", "text/csv", category="calculation")
    assert doc["status"] in {"parsed", "analyzed"}  # "analyzed" if GEMINI_API_KEY happens to be configured

    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert len(chunks) == 1
    assert "33" in chunks[0]["text"]


def test_csv_with_title_line_before_header_is_not_mistaken_for_header(tmp_path):
    # real project exports often prefix the table with a single-cell title/rev line —
    # the header detector must skip it, not treat it as the header or a fake data row.
    csv_path = tmp_path / "Equipment_List.csv"
    csv_path.write_text(
        "OO 프로젝트 - 설비 목록 / DOC-001 Rev.A\n"
        "TAG,Manufacturer,Power\n"
        "M-201,Siemens,15\n",
        encoding="utf-8",
    )

    project_id = _new_project()
    doc = _upload(project_id, csv_path, "Equipment_List.csv", "text/csv", category="equipment_list")
    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()

    assert len(chunks) == 1  # only the real data row — title line and header row are context, not chunks
    assert chunks[0]["location"] == "Row3"
    assert "TAG" in chunks[0]["text"]  # real header carried as context
    assert "OO 프로젝트" in chunks[0]["text"]  # title line also kept, just not mistaken for the header
    assert "M-201" in chunks[0]["text"] and "Siemens" in chunks[0]["text"]


def test_csv_multi_row_gets_one_chunk_per_row_with_header_context(tmp_path):
    csv_path = tmp_path / "Equipment_List.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["TAG", "Power"])
        writer.writerow(["M-201", "15"])
        writer.writerow(["M-202", "30"])

    project_id = _new_project()
    doc = _upload(project_id, csv_path, "Equipment_List.csv", "text/csv", category="equipment_list")
    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()

    # one chunk per data row (2 rows -> 2 chunks), each a distinct, traceable location
    assert {c["location"] for c in chunks} == {"Row2", "Row3"}
    row2 = next(c for c in chunks if c["location"] == "Row2")
    # header carried along for column context, but the other row's value isn't in this chunk
    assert "TAG" in row2["text"] and "M-201" in row2["text"] and "M-202" not in row2["text"]


def test_xlsx_multi_sheet_locations_are_prefixed_by_sheet(tmp_path):
    import openpyxl as _oxl

    xlsx_path = tmp_path / "Two_Sheets.xlsx"
    wb = _oxl.Workbook()
    ws1 = wb.active
    ws1.title = "Motors"
    ws1.append(["TAG", "Power"])
    ws1.append(["M-201", 15])
    ws2 = wb.create_sheet("Pumps")
    ws2.append(["TAG", "Flow"])
    ws2.append(["P-201", 120])
    wb.save(xlsx_path)

    project_id = _new_project()
    doc = _upload(
        project_id, xlsx_path, "Two_Sheets.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        category="equipment_list",
    )
    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert {c["location"] for c in chunks} == {"Motors!Row2", "Pumps!Row2"}


def test_txt_parsing(tmp_path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("TAG M-201 power rating is 30 kW.", encoding="utf-8")

    project_id = _new_project()
    doc = _upload(project_id, txt_path, "notes.txt", "text/plain")
    assert doc["status"] in {"parsed", "analyzed"}  # "analyzed" if GEMINI_API_KEY happens to be configured

    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert len(chunks) == 1
    assert "30 kW" in chunks[0]["text"]


def test_image_upload_has_no_chunks_but_succeeds(tmp_path):
    img_path = tmp_path / "nameplate.png"
    # minimal 1x1 PNG
    img_path.write_bytes(bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000a49444154789c6360000002000100"
        "5e2707a30000000049454e44ae426082"
    ))

    project_id = _new_project()
    doc = _upload(project_id, img_path, "nameplate.png", "image/png", category="drawing")
    assert doc["status"] in {"uploaded", "analyzed"}  # no deterministic parser for images either way

    chunks = client.get(f"/api/projects/{project_id}/documents/{doc['id']}/content").json()
    assert chunks == []
